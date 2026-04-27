"""Generate the MCP SF Value Proposition Word document."""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# -- Style setup --
style = doc.styles["Normal"]
font = style.font
font.name = "Calibri"
font.size = Pt(11)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    h_style = doc.styles[f"Heading {level}"]
    h_style.font.name = "Calibri"
    h_style.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    if level == 1:
        h_style.font.size = Pt(24)
        h_style.paragraph_format.space_before = Pt(24)
        h_style.paragraph_format.space_after = Pt(12)
    elif level == 2:
        h_style.font.size = Pt(18)
        h_style.paragraph_format.space_before = Pt(18)
        h_style.paragraph_format.space_after = Pt(8)
    else:
        h_style.font.size = Pt(14)
        h_style.paragraph_format.space_before = Pt(12)
        h_style.paragraph_format.space_after = Pt(6)

# Margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def add_quote(text):
    """Add a styled quote block."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(f"\u201c{text}\u201d")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)


def add_tools_line(text):
    """Add a tools reference line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run("Herramientas: ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run2 = p.add_run(text)
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_separator():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\u2500" * 40)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(8)


def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


# ============================================================
# COVER / TITLE
# ============================================================

# Add some spacing before title
for _ in range(4):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("SF Assistant MCP")
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
run.bold = True

subtitle_p = doc.add_paragraph()
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle_p.add_run("El copiloto inteligente para consultoría SuccessFactors")
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

tagline_p = doc.add_paragraph()
tagline_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = tagline_p.add_run("Seguridad | Productividad | Inteligencia")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
run.italic = True

doc.add_page_break()

# ============================================================
# 1. LA HISTORIA
# ============================================================

doc.add_heading("La historia empieza con una pregunta incómoda", level=1)

doc.add_paragraph(
    "Son las 11 de la noche. Un consultor senior de SuccessFactors tiene abierto "
    "Provisioning en una pestaña, un Excel con 4,000 registros de EmpJob en otra, "
    "y la API de OData en una tercera. Está preparando un cutover para producción "
    "que se ejecuta en 48 horas. Necesita verificar que las business rules del "
    "entorno QA coinciden con DEV, que los picklists de 12 países están completos, "
    "y que ningún registro de compensación tiene gaps en effective dating."
)

doc.add_paragraph(
    "Revisa campo por campo. Copia y pega entre pestañas. Hace queries manuales en "
    "Postman. Son las 2 de la mañana cuando termina, y todavía no está seguro de "
    "que no se le haya escapado nada."
)

p = doc.add_paragraph()
run = p.add_run(
    "Esta es la realidad de la consultoría SF hoy. "
)
run.bold = True
run2 = p.add_run(
    "Trabajo repetitivo, propenso a errores humanos, y con un costo invisible: "
    "el tiempo que no se dedica a pensar, diseñar y asesorar al cliente."
)

doc.add_paragraph(
    "SF Assistant MCP existe para cambiar esa ecuación. Pero antes de hablar de "
    "productividad, hablemos de lo que realmente importa."
)

add_separator()

# ============================================================
# 2. SEGURIDAD
# ============================================================

doc.add_heading("Primero lo primero: seguridad", level=1)

doc.add_paragraph(
    "Cuando una herramienta se conecta a un sistema que almacena datos de nómina, "
    "información personal protegida, y estructuras organizacionales completas de "
    "una empresa, la pregunta no es \u201c¿qué puede hacer?\u201d sino "
    "\u201c¿qué tan seguro es lo que hace?\u201d"
)

# --- OIDC ---
doc.add_heading("Autenticación: OIDC de dos pasos vía IAS", level=2)

doc.add_paragraph(
    "SF Assistant MCP no almacena ni transmite credenciales en texto plano. Utiliza "
    "el flujo estándar de SAP para autenticación empresarial:"
)

p1 = doc.add_paragraph(style="List Bullet")
run = p1.add_run("Paso 1 \u2014 Password grant a IAS: ")
run.bold = True
p1.add_run(
    "Se obtiene un id_token del Identity Authentication Service de SAP, el mismo "
    "servicio que autentica a los usuarios en la plataforma BTP."
)

p2 = doc.add_paragraph(style="List Bullet")
run = p2.add_run("Paso 2 \u2014 JWT-Bearer exchange: ")
run.bold = True
p2.add_run(
    "Ese id_token se intercambia por un access_token específico de SuccessFactors, "
    "con alcance limitado a la instancia configurada."
)

doc.add_paragraph(
    "No hay contraseñas viajando en cada request. No hay Basic Auth. No hay tokens "
    "hardcodeados. El flujo es el mismo que SAP recomienda para integraciones de producción."
)

# --- Tokens ---
doc.add_heading("Tokens con ciclo de vida controlado", level=2)

doc.add_paragraph(
    "Los tokens de acceso se cachean en memoria con un buffer de seguridad de 60 "
    "segundos antes de su expiración. Cuando expiran, se renueva el ciclo completo. "
    "No hay tokens persistidos en disco, no hay refresh tokens almacenados en "
    "archivos de configuración."
)

# --- Credentials ---
doc.add_heading("Credenciales: nunca en el código", level=2)

doc.add_paragraph(
    "Las credenciales se resuelven desde dos fuentes posibles, nunca desde el código fuente:"
)

p = doc.add_paragraph(style="List Bullet")
run = p.add_run("Variables de entorno (.env): ")
run.bold = True
p.add_run(
    "Para desarrollo local y entornos controlados. El archivo .env está excluido "
    "del repositorio por diseño."
)

p = doc.add_paragraph(style="List Bullet")
run = p.add_run("Google Cloud Secret Manager: ")
run.bold = True
p.add_run(
    "Para entornos productivos y multi-tenant. Las credenciales se almacenan "
    "cifradas en GCP y se recuperan bajo demanda, siguiendo el principio de mínimo privilegio."
)

doc.add_paragraph(
    "La arquitectura de credential store es pluggable. Si mañana la política de "
    "seguridad exige usar Azure Key Vault o AWS Secrets Manager, el patrón de "
    "abstracción ya está preparado."
)

# --- OData injection ---
doc.add_heading("Protección contra inyección OData", level=2)

doc.add_paragraph(
    "Cada valor proporcionado por el usuario pasa por sanitización antes de "
    "incluirse en queries OData. Los nombres de campo se validan contra un patrón "
    "estricto (solo alfanuméricos y underscores, máximo 128 caracteres). Los "
    "valores de string se escapan para prevenir inyección de filtros. No se "
    "construyen queries concatenando strings sin validar."
)

# --- PII Protection ---
doc.add_heading("Protección de datos personales (PII): la barrera más importante", level=2)

doc.add_paragraph(
    "Imagine la siguiente escena: un consultor, preparando una migración de datos, "
    "escribe una consulta natural en el asistente: \"dame los empleados de México "
    "con su información personal\". Sin protección, el sistema ejecuta esa consulta "
    "literalmente y devuelve, en segundos, los números de seguro social de 5,000 "
    "empleados directamente en el chat. Un screenshot accidental, un log mal "
    "configurado, un pantallón compartido en una videollamada — y hay una brecha "
    "de datos. Con consecuencias legales reales."
)

doc.add_paragraph(
    "Este escenario no es hipotético. Es el riesgo exacto que el MCP elimina con "
    "una capa de protección PII construida sobre la política de clasificación de "
    "datos de Veritas Prime."
)

# PII tier table
doc.add_paragraph()
pii_table = doc.add_table(rows=5, cols=3)
pii_table.alignment = WD_TABLE_ALIGNMENT.CENTER
pii_table.style = "Table Grid"

pii_headers = ["Nivel", "Ejemplos en SuccessFactors", "Acción del sistema"]
for i, h in enumerate(pii_headers):
    cell = pii_table.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10)
    set_cell_shading(cell, "1A3C6E")

pii_rows = [
    ("ALTAMENTE CONFIDENCIAL", "SSN, número de pasaporte, cuenta bancaria, datos de salud", "Bloqueado — la consulta no se ejecuta"),
    ("CONFIDENCIAL", "Salario, historial disciplinario, documentos de identidad", "Advertencia explícita con registro"),
    ("INTERNO", "Dirección, fecha de nacimiento, género", "Permitido con monitoreo"),
    ("PÚBLICO", "Nombre, email corporativo, cargo, departamento", "Acceso normal"),
]

for idx, (level, examples, action) in enumerate(pii_rows, start=1):
    pii_table.rows[idx].cells[0].text = level
    pii_table.rows[idx].cells[1].text = examples
    pii_table.rows[idx].cells[2].text = action
    for cell in pii_table.rows[idx].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
        if idx % 2 == 0:
            set_cell_shading(cell, "F0F4FA")

doc.add_paragraph()

doc.add_paragraph(
    "La protección opera en dos capas independientes — defensa en profundidad:"
)

p1 = doc.add_paragraph(style="List Bullet")
p1.add_run("Capa 1 — Validación pre-consulta: ").bold = True
p1.add_run(
    "antes de ejecutar cualquier query contra SuccessFactors, el sistema verifica "
    "la entidad solicitada y los campos en $select. Si alguno pertenece al nivel "
    "ALTAMENTE CONFIDENCIAL, la consulta se bloquea antes de salir del servidor. "
    "La entidad PerNationalId (que contiene SSNs y números de documento) está "
    "completamente bloqueada — no puede ser consultada por ninguna herramienta."
)

p2 = doc.add_paragraph(style="List Bullet")
p2.add_run("Capa 2 — Limpieza post-respuesta: ").bold = True
p2.add_run(
    "incluso si una respuesta de SuccessFactors incluye campos sensibles de manera "
    "inesperada (por ejemplo, en un expand implícito), el sistema los elimina de "
    "la respuesta antes de devolvérsela al usuario. La segunda capa garantiza que "
    "aunque la primera falle, los datos no llegan al chat."
)

doc.add_paragraph(
    "Una regla adicional protege contra la correlación indirecta: si una consulta "
    "combina 3 o más campos de identificación indirecta (fecha de nacimiento, "
    "género, código postal, origen étnico), el sistema escala la advertencia. "
    "Porque aunque cada campo sea inocente por separado, combinados pueden "
    "re-identificar a una persona de forma única."
)

doc.add_paragraph(
    "Cuando una consulta es bloqueada, el sistema no devuelve un error críptico. "
    "Explica exactamente por qué fue bloqueada, qué clasificación tienen los "
    "campos solicitados, y sugiere una alternativa segura con los campos permitidos. "
    "El consultor entiende la razón y puede replantear su consulta correctamente."
)

# --- Dry run ---
doc.add_heading("Operaciones de escritura: seguras por defecto", level=2)

doc.add_paragraph(
    "La herramienta execute_upsert opera en modo dry_run por defecto. Esto "
    "significa que cuando un consultor genera una operación de importación, el "
    "sistema simula la ejecución y muestra qué haría, sin tocar un solo registro "
    "en la instancia. La escritura real requiere una confirmación explícita. "
    "Un error de dedo no se convierte en un incidente de producción."
)

# --- Multi-tenant ---
doc.add_heading("Multi-tenant con aislamiento", level=2)

doc.add_paragraph(
    "Cada tenant tiene su propio cache de tokens, keyed por tenant_id. Un consultor "
    "que trabaja con múltiples clientes simultáneamente no corre riesgo de que una "
    "query a la instancia del Cliente A se ejecute accidentalmente con las "
    "credenciales del Cliente B. El aislamiento es estructural, no por convención."
)

# --- Rate limiting ---
doc.add_heading("Rate limiting integrado", level=2)

doc.add_paragraph(
    "Todas las llamadas a la API de SuccessFactors pasan por un semáforo asíncrono "
    "que limita a 5 requests concurrentes. Esto protege la instancia del cliente "
    "de sobrecarga, respeta los límites de throttling de SAP, y evita que un "
    "análisis masivo degrade el rendimiento del sistema para los usuarios finales."
)

# --- Read-only audits ---
doc.add_heading("Todas las auditorías son read-only", level=2)

doc.add_paragraph(
    "Las herramientas de auditoría, detección de anomalías, trazado de reglas y "
    "verificación de permisos operan exclusivamente en modo lectura. Diagnostican "
    "sin modificar. Un consultor puede investigar un problema de datos a fondo sin "
    "riesgo de alterar el estado del sistema."
)

# -- Security summary table --
doc.add_paragraph()
table = doc.add_table(rows=9, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

headers = ["Aspecto de seguridad", "Implementación"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10)
    set_cell_shading(cell, "1A3C6E")

rows_data = [
    ("Autenticación", "OIDC 2-step via SAP IAS (id_token + JWT-bearer)"),
    ("Tokens", "En memoria, auto-renovación con buffer de 60s"),
    ("Credenciales", "Env vars o GCP Secret Manager (pluggable)"),
    ("Inyección OData", "Sanitización de valores + validación de nombres de campo"),
    ("Escritura", "dry_run por defecto en todas las operaciones de upsert"),
    ("Multi-tenant", "Aislamiento por tenant_id en cache de tokens"),
    ("Rate limiting", "Semáforo asíncrono: máx. 5 requests concurrentes"),
    ("Protección PII", "Clasificación 4 niveles, bloqueo pre-consulta + limpieza post-respuesta, regla de combinación indirecta"),
]

for idx, (aspect, impl) in enumerate(rows_data, start=1):
    table.rows[idx].cells[0].text = aspect
    table.rows[idx].cells[1].text = impl
    for cell in table.rows[idx].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)
        if idx % 2 == 0:
            set_cell_shading(cell, "F0F4FA")

doc.add_page_break()

# ============================================================
# 3. EL PROBLEMA
# ============================================================

doc.add_heading("El problema real de la consultoría SF", level=1)

doc.add_paragraph(
    "Un proyecto típico de implementación de SuccessFactors tiene fases bien "
    "conocidas: discovery, diseño, configuración, migración de datos, testing, "
    "cutover, go-live, y soporte post-go-live. En cada una de estas fases, los "
    "consultores invierten una proporción significativa de su tiempo en tareas "
    "que son necesarias pero no son estratégicas:"
)

pain_points = [
    "Explorar metadatos para entender qué campos existen, cuáles son editables, cuáles requieren picklists.",
    "Construir queries OData manualmente para extraer datos de validación.",
    "Comparar configuraciones entre entornos (DEV, QA, PROD) copiando y pegando.",
    "Generar templates de importación revisando documentación de la API.",
    "Escribir documentación funcional que describe lo que ya se configuró.",
    "Validar datos de migración registro por registro.",
    "Preparar checklists de cutover desde cero para cada proyecto.",
]

for point in pain_points:
    doc.add_paragraph(point, style="List Bullet")

p = doc.add_paragraph()
run = p.add_run(
    "Cada una de estas tareas es importante. Ninguna requiere el juicio experto "
    "de un consultor senior. Y todas consumen horas que podrían dedicarse a "
    "diseñar soluciones, asesorar al cliente, o anticipar problemas."
)
run.italic = True

add_separator()

# ============================================================
# 4. LAS 50 HERRAMIENTAS
# ============================================================

doc.add_heading("50 herramientas, organizadas por cómo trabajas", level=1)

doc.add_paragraph(
    "SF Assistant MCP no es un dashboard ni una interfaz gráfica. Es un conjunto "
    "de 50 herramientas especializadas que se integran directamente en el flujo "
    "de trabajo del consultor a través del Model Context Protocol. Esto significa "
    "que se usan conversacionalmente, como si le pidieras algo a un colega que "
    "conoce la API de SuccessFactors de memoria."
)

# --- Discovery ---
doc.add_heading("Discovery: entiende la instancia en segundos", level=2)

add_quote("Muéstrame todos los campos de EmpJob que son editables y tienen picklist.")

doc.add_paragraph(
    "En lugar de navegar por $metadata XML o buscar en la documentación de SAP, "
    "el consultor pregunta y obtiene una respuesta estructurada con los campos, "
    "sus tipos, labels, y anotaciones SAP (creatable, updatable, filterable)."
)

add_tools_line("get_entity_metadata, list_entities, get_picklist_values, query_odata")

# --- Business Rules ---
doc.add_heading("Business Rules: del requerimiento al test case", level=2)

add_quote(
    "Necesito una regla que cuando el país sea México y el event reason sea "
    "'Promoción', setee automáticamente el campo custom de bono al 15%."
)

doc.add_paragraph(
    "El flujo tradicional: el consultor interpreta el requerimiento, busca qué "
    "campos existen, verifica que los valores del picklist sean correctos, escribe "
    "la especificación, configura la regla, documenta lo que hizo, y crea casos de "
    "prueba. Con MCP, ese flujo se comprime:"
)

steps = [
    ("validate_rule_fields", "confirma que los campos y valores existen en la instancia."),
    ("generate_rule_spec", "genera la especificación formal de la regla desde lenguaje natural."),
    ("generate_rule_doc", "produce la guía de implementación paso a paso."),
    ("generate_rule_test", "crea la matriz de casos de prueba con escenarios positivos, negativos y edge cases."),
]

for i, (tool, desc) in enumerate(steps, 1):
    p = doc.add_paragraph()
    run = p.add_run(f"{i}. {tool}")
    run.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    p.add_run(f" \u2014 {desc}")

p = doc.add_paragraph()
run = p.add_run("Lo que antes tomaba medio día, ahora toma minutos.")
run.bold = True
p.add_run(
    " Y con menos margen de error, porque la validación contra la instancia "
    "real ocurre desde el primer paso."
)

# --- Employee ---
doc.add_heading("Datos de empleado: respuestas inmediatas", level=2)

add_quote("Compárame el perfil de estos dos empleados que deberían tener la misma configuración.")

doc.add_paragraph(
    "Las herramientas de empleado permiten búsqueda en lenguaje natural, perfiles "
    "completos, organigramas, historial de empleo, y comparación lado a lado. Lo "
    "que normalmente requiere múltiples queries OData manuales se resuelve en una "
    "sola interacción."
)

add_tools_line("search_employees, get_employee_profile, get_org_chart, get_employee_history, compare_employees")

# --- Instance Compare ---
doc.add_heading("Comparación de instancias: DEV vs QA vs PROD sin dolor", level=2)

add_quote("Compárame las picklists de departamento entre QA y producción.")

doc.add_paragraph(
    "Antes de cada transporte o cutover, los consultores necesitan verificar que "
    "la configuración entre entornos es consistente. Hacerlo manualmente es tedioso "
    "y propenso a omisiones. Las herramientas de comparación extraen la configuración "
    "de ambas instancias y muestran exactamente qué difiere."
)

add_tools_line("compare_instance_config, compare_picklists, compare_business_rules, compare_foundation_objects")

# --- Audit ---
doc.add_heading("Auditoría y troubleshooting: diagnostica sin riesgo", level=2)

add_quote("Audita los datos de este empleado que reporta un error en su recibo de nómina.")

doc.add_paragraph(
    "Las herramientas de auditoría verifican integridad de datos, detectan anomalías, "
    "trazan ejecución de reglas, validan permisos, y encuentran gaps en effective "
    "dating. Todo en modo lectura, todo contra la instancia real."
)

add_tools_line("audit_employee_data, find_data_anomalies, trace_rule_execution, check_permission_access, validate_effective_dating")

# --- Migration ---
doc.add_heading("Migración de datos: el proceso más crítico, bajo control", level=2)

add_quote("Valida este archivo de migración de EmpJob antes de cargarlo.")

doc.add_paragraph(
    "La migración de datos es donde más incidentes ocurren en proyectos SF. Un "
    "campo mal mapeado, un formato de fecha incorrecto, o un registro duplicado "
    "pueden causar horas de retrabajo. Las herramientas de migración validan "
    "archivos contra metadatos reales, calculan el orden correcto de carga "
    "(respetando dependencias entre entidades), reconcilian datos post-migración, "
    "y generan archivos de purge cuando algo sale mal."
)

add_tools_line("validate_migration_file, generate_migration_sequence, reconcile_data, generate_purge_file")

# --- Documentation ---
doc.add_heading("Documentación: generada, no escrita", level=2)

add_quote("Genera el functional spec de la configuración actual de Employee Central.")

doc.add_paragraph(
    "La documentación es el eterno pendiente de todo proyecto. Con MCP, los specs "
    "funcionales, checklists de cutover, diccionarios de datos, y scripts de prueba "
    "se generan a partir de la configuración real de la instancia. No se inventan: "
    "se documentan los hechos."
)

add_tools_line("generate_functional_spec, generate_cutover_checklist, generate_data_dictionary, generate_test_script")

# --- Go-Live ---
doc.add_heading("Go-Live: validación antes de cruzar el punto de no retorno", level=2)

add_quote("Ejecuta los checks de go-live en la instancia de producción.")

doc.add_paragraph(
    "El go-live es el momento de mayor presión en un proyecto SF. Las herramientas "
    "de validación pre-go-live ejecutan checks críticos, importantes, y recomendados "
    "contra la instancia, y generan reportes de reconciliación que documentan el "
    "estado exacto del sistema en el momento del corte."
)

add_tools_line("run_go_live_checks, generate_reconciliation_report")

# --- Country ---
doc.add_heading("Soporte multi-país: LATAM y más allá", level=2)

add_quote(
    "Dame los campos específicos de México y valida que la configuración cumple "
    "con los requisitos locales."
)

doc.add_paragraph(
    "Cada país tiene campos obligatorios, formatos específicos, y requisitos de "
    "compliance propios. Las herramientas de país identifican esos campos, validan "
    "cumplimiento, y generan guías de configuración específicas por país."
)

add_tools_line("get_country_specific_fields, validate_country_compliance, generate_country_config_guide")

doc.add_page_break()

# ============================================================
# 5. EL CAMBIO REAL
# ============================================================

doc.add_heading("El cambio real: de ejecutor a asesor", level=1)

doc.add_paragraph(
    "La propuesta de valor de SF Assistant MCP no se mide solo en horas ahorradas, "
    "aunque las ahorra. Se mide en qué tipo de consultor puedes ser cuando las "
    "tareas mecánicas dejan de consumir tu día."
)

# Comparison table
table = doc.add_table(rows=7, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

# Header
for i, h in enumerate(["Antes", "Después"]):
    cell = table.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(11)
    set_cell_shading(cell, "1A3C6E")

comparison = [
    ("2 horas explorando metadatos", "2 minutos preguntando"),
    ("Medio día escribiendo specs de reglas", "15 minutos validando lo generado"),
    ("1 día comparando entornos antes del cutover", "10 minutos con comparación automatizada"),
    ("Documentación atrasada o incompleta", "Documentación generada al instante"),
    ("Validación manual de migración", "Validación automática contra metadatos reales"),
    ("Go-live con checklist manual", "Go-live con validación sistemática"),
]

for idx, (before, after) in enumerate(comparison, start=1):
    table.rows[idx].cells[0].text = before
    table.rows[idx].cells[1].text = after
    for cell in table.rows[idx].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)
        if idx % 2 == 0:
            set_cell_shading(cell, "F0F4FA")

doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run(
    "El consultor deja de ser la persona que ejecuta tareas repetitivas y se "
    "convierte en la persona que piensa, diseña y decide. "
)
run.bold = True
p.add_run(
    "La herramienta se encarga de lo mecánico. El consultor se encarga de lo que "
    "importa: entender el negocio del cliente, diseñar la solución correcta, y "
    "asegurar que la implementación refleje la estrategia de RRHH de la organización."
)

add_separator()

# ============================================================
# 6. PARA QUIÉN
# ============================================================

doc.add_heading("Para quién es esto", level=1)

audiences = [
    ("Consultores SF", "que implementan Employee Central, compensación, o cualquier módulo y quieren reducir el trabajo manual sin sacrificar calidad."),
    ("Leads técnicos", "que necesitan comparar entornos, validar migraciones, y preparar go-lives con confianza."),
    ("Equipos de soporte", "que diagnostican incidentes y necesitan datos precisos rápidamente, sin riesgo de modificar la instancia."),
    ("Firmas de consultoría", "que buscan escalar su capacidad de delivery sin escalar proporcionalmente su headcount."),
]

for title, desc in audiences:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(f"{title} ")
    run.bold = True
    p.add_run(desc)

add_separator()

# ============================================================
# 7. CONCLUSIÓN
# ============================================================

doc.add_heading("Conclusión", level=1)

doc.add_paragraph(
    "La consultoría de SuccessFactors siempre ha requerido conocimiento profundo "
    "del sistema. Eso no cambia. Lo que cambia es que ese conocimiento ya no tiene "
    "que gastarse en tareas que una máquina puede hacer mejor, más rápido, y con "
    "menos errores."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
run = p.add_run(
    "SF Assistant MCP es seguro por diseño, eficiente por arquitectura, "
    "y útil por vocación.\nNo reemplaza al consultor. Le devuelve el tiempo para ser consultor."
)
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

# Save
output_path = "docs/SF_Assistant_MCP_Value_Proposition.docx"
doc.save(output_path)
print(f"Document saved to: {output_path}")
