"""
Template parser for CSV/Excel import files.

Parses import files into structured headers + rows, and validates
headers against SF entity metadata.
"""

from __future__ import annotations

import csv
import io
from typing import Optional


def parse_csv_template(csv_content: str) -> dict:
    """Parse CSV content into headers and rows.

    Args:
        csv_content: Raw CSV string

    Returns:
        dict with 'headers' (list[str]) and 'rows' (list[dict])
    """
    reader = csv.DictReader(io.StringIO(csv_content))
    headers = reader.fieldnames or []
    rows = [row for row in reader]
    return {
        "headers": list(headers),
        "rows": rows,
        "row_count": len(rows),
    }


def parse_excel_template(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
) -> dict:
    """Parse Excel file bytes into headers and rows.

    Args:
        file_bytes: Raw bytes of an .xlsx file
        sheet_name: Optional sheet name. Default: active sheet.

    Returns:
        dict with 'headers' (list[str]) and 'rows' (list[dict])
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"error": "openpyxl is required for Excel parsing. Install with: uv add openpyxl"}

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {"headers": [], "rows": [], "row_count": 0}

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    rows = []
    for row in rows_iter:
        row_dict = {}
        for i, cell in enumerate(row):
            if i < len(headers) and headers[i]:
                row_dict[headers[i]] = cell
        if any(v is not None for v in row_dict.values()):
            rows.append(row_dict)

    wb.close()
    return {
        "headers": headers,
        "rows": rows,
        "row_count": len(rows),
    }


def validate_template_against_metadata(
    headers: list[str],
    entity_metadata: dict,
) -> dict:
    """Validate import headers against entity metadata.

    Args:
        headers: List of column headers from the template
        entity_metadata: Parsed entity metadata (from metadata_parser)

    Returns:
        Validation report with field-level results
    """
    properties = entity_metadata.get("properties", {})
    keys = entity_metadata.get("keys", [])

    field_results = []
    valid_count = 0
    invalid_count = 0

    for header in headers:
        if header in properties:
            prop = properties[header]
            field_results.append({
                "header": header,
                "mapped_field": header,
                "status": "valid",
                "field_type": prop.get("type", ""),
                "is_key": header in keys,
                "is_required": prop.get("required", False),
                "notes": "",
            })
            valid_count += 1
        else:
            field_results.append({
                "header": header,
                "mapped_field": "",
                "status": "invalid",
                "field_type": "",
                "is_key": False,
                "is_required": False,
                "notes": f"Field '{header}' not found in entity metadata",
            })
            invalid_count += 1

    # Check for missing required/key fields
    missing_keys = [k for k in keys if k not in headers]
    warnings = []
    if missing_keys:
        warnings.append(f"Missing key fields: {', '.join(missing_keys)}")

    required_fields = [
        name for name, prop in properties.items()
        if prop.get("required") and name not in headers
    ]
    if required_fields:
        warnings.append(f"Missing required fields: {', '.join(required_fields)}")

    return {
        "valid": invalid_count == 0 and not missing_keys,
        "total_headers": len(headers),
        "valid_headers": valid_count,
        "invalid_headers": invalid_count,
        "field_results": field_results,
        "missing_key_fields": missing_keys,
        "warnings": warnings,
    }
