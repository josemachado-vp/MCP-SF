from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

all_users = [
    ("BSIADMIN","Stephanie","User","N/A","N/A","N/A","","dummy@sap.com","BSIADMIN","t"),
    ("SuperAdmin","Jimena","Johnston","N/A","N/A","N/A","","dummy@sap.com","SuperAdmin","t"),
    ("sfadmin","Michael","Kumar","Corporate Services (CORP_SVCS)","Global Human Resources (50007726)","Corporate - US-Philadelphia (1710-2001)","Administrative Support","dummy@sap.com","sfadmin","t"),
    ("sfapiuserECP","William","Admin","N/A","N/A","N/A","","dummy@sap.com","sfapiuserECP","t"),
    ("superadmin","SAPTest!SAPTestÑ","Armijo","N/A","N/A","N/A","","dummy@sap.com","superadmin","t"),
    ("VPadminLS","Ted","Cook","VPUK","VPUK","N/A","","dummy@sap.com","VPadminLS","t"),
    ("SFADMIN","David","Tello","N/A","N/A","N/A","","dummy@sap.com","SFADMIN","t"),
    ("VPadminReport","Super","Test","N/A","N/A","N/A","","dummy@sap.com","VPadminReport","t"),
    ("VPadminJN","Jacob","","N/A","N/A","N/A","","dummy@sap.com","VPadminJN","t"),
    ("HRADMIN","Marge","Hukrwza","Information Technology (IT)","IT Business Office (5000137)","Boston (10000-0001)","HR Admin","dummy@sap.com","HRADMIN","t"),
    ("lhradmin","Janet","Stephens","Corporate Services (CORP_SVCS)","N/A","Corporate - US-Philadelphia (0300-0001)","HR Admin","dummy@sap.com","lhradmin","t"),
    ("VPadminIntegration","Emma","Garrity","N/A","N/A","N/A","","dummy@sap.com","VPadminIntegration","t"),
    ("PSADMIN","Andrea","Bartex","N/A","N/A","N/A","","dummy@sap.com","PSADMIN","t"),
    ("SYSADMIN","Michael","Test","Information Technology (IT)","IT Business Office (5000137)","Boston (10000-0001)","System Admin","dummy@sap.com","SYSADMIN","t"),
    ("IPSADMIN_OPHXMPP","Doug","Martin","N/A","N/A","N/A","","dummy@sap.com","IPSADMIN_OPHXMPP","t"),
    ("VPadminMR","Joanne","Atchison","N/A","N/A","N/A","","dummy@sap.com","VPadminMR","t"),
    ("adminOPP","AdminOptimusHXM","Van","N/A","N/A","N/A","","dummy@sap.com","adminOPP","t"),
    ("VPadminTest","Rodrigo","Peña","N/A","N/A","N/A","","dummy@sap.com","VPadminTest","t"),
    ("superadmin1","Ian","Villanueva","N/A","N/A","N/A","","dummy@sap.com","superadmin1","t"),
    ("VPadminGL","LOA01","Santos","N/A","N/A","N/A","","dummy@sap.com","VPadminGL","t"),
    ("CompAdmin","Nettie","Golden","N/A","N/A","N/A","","dummy@sap.com","CompAdmin","t"),
    ("10000255","Gamora","Guardian","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","Finance Manager","noemail@company.com","10000255","t"),
    ("10000258","RetestThursday","DataFlow","Corporate (30000005)","Corporate Administrative (40000021)","Atlanta (1710-2015)","Finance Manager","noemail@company.com","10000258","t"),
    ("10000260","Mandy","Moore","Corporate (30000005)","Corporate Administrative (40000021)","Atlanta (1710-2015)","Finance Manager","noemail@company.com","10000260","t"),
    ("10000261","Warren","Worthington","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","Finance Manager","noemail@company.com","10000261","t"),
    ("10000263","Reed","Richards","Chile Operations (30000024)","Operations and Logistics CHL (40000023)","Santiago (CL01_1001)","Supply Chain Coordinator CHL","noemail@company.com","10000263","t"),
    ("10000266","Hobby","Lobby","Corporate (30000005)","Corporate Administrative (40000021)","Atlanta (1710-2015)","Finance Manager","noemail@company.com","10000266","t"),
    ("172","Sailor","Moon","Finance (30000006)","Finance (40000009)","Atlanta (1710-2015)","Finance Manager","noemail@company.com","172","t"),
    ("200002","Tony","Stark","Logistics (30000016)","Business Operations (40000015)","Atlanta (1710-2015)","COO","","200002","t"),
    ("200006","Thor","Odinson","Logistics (30000016)","Business Operations (40000015)","Atlanta (1710-2015)","Business Director","","200006","t"),
    ("200008","Wanda","Maximoff","Corporate (30000005)","Corporate Administrative (40000021)","Ciudad de Mexico DF (MX_1001)","Administrative Support","","200008","t"),
    ("200009","Scott","Lang","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","Administrative Support","","200009","t"),
    ("200011","Stephen","Strange","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","CA CEO","","200011","t"),
    ("200013","T'Challa","Panther","Corporate (30000005)","Corporate Administrative (40000021)","Ciudad de Mexico DF (MX_1001)","Executive Management","","200013","t"),
    ("200014","Victor","Von Doom","Corporate (30000005)","Corporate Administrative (40000021)","Lima - Oficina Central (PE01_1001)","Administrative Support","","200014","t"),
    ("200015","Peter","Quill","Global HR (30000014)","Human Resources (40000002)","Atlanta (1710-2015)","CHRO","","200015","t"),
    ("200016","Drax","Destroyer","Global HR (30000014)","Human Resources (40000002)","Atlanta (1710-2015)","HR Director","","200016","t"),
    ("200017","Rocket","Raccoon","Global HR (30000014)","Human Resources (40000002)","Atlanta (1710-2015)","HRBP","","200017","t"),
    ("200018","Groot","Tree","Global HR (30000014)","Human Resources (40000002)","Atlanta (1710-2015)","HRBP","","200018","t"),
    ("200023","Nick","Fury","Corporate (30000005)","Corporate Administrative (40000021)","Atlanta (1710-2015)","Chief Executive Officer","","200023","t"),
    ("200024","Matt","Murdock","Information Technology (30000009)","Information Technology (40000011)","Atlanta (1710-2015)","IT Analyst","","200024","t"),
    ("200025","Wade","Wilson","Information Technology (30000009)","Information Technology (40000011)","Atlanta (1710-2015)","IT Analyst","","200025","t"),
    ("10000259","Monica","Rambeau","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","Finance Manager","noemail@company.com","10000259","t"),
    ("10000262","Alexis","Sánchez","Chile Operations (30000024)","Operations and Logistics CHL (40000023)","Santiago (CL01_1001)","Operations and Logistics CHL","noemail@company.com","10000262","t"),
    ("10000267","Ana","Banana","Chile Operations (30000024)","Operations and Logistics CHL (40000023)","Santiago (CL01_1001)","Supply Chain Coordinator CHL","","10000267","t"),
    ("10000269","Kiyotaka","Ayanokoji","Manufacturing (30000001)","Production US (40000003)","San Francisco (1710-2003)","Management & Planning","test.kiyotaka@veritasprime.com","10000269","t"),
    ("PAY_PORTAL_API_USER","Pay Portal","API","N/A","N/A","N/A","","dummy@successfactors.com","PAY_PORTAL_API_USER","t"),
    ("VP_CONNECT_API_USER","VP","Connect API","N/A","N/A","N/A","","dummy@successfactors.com","VP_CONNECT_API_USER","t"),
    ("adminSPA","Sankar","Palamadai","N/A","N/A","N/A","","sankara.palamadai@veritasprime.com","adminSPA","t"),
    ("adminKK","Ken","Kilmer","N/A","N/A","N/A","","kenneth.kilmer@veritasprime.com","adminKK","t"),
    ("adminFH","Frank","Huber","N/A","N/A","N/A","","frank.huber@veritasprime.com","adminFH","t"),
    ("adminCN2","Chris","Nguyen","N/A","N/A","N/A","","chris.nguyen@veritasprime.com","adminCN2","t"),
    ("adminZL","Admin","Ziplyne","N/A","N/A","N/A","","kayla.toelle+2@veritasprime.com","adminZL","t"),
    ("IPSADMIN","Janet","Roth","N/A","N/A","N/A","","dummy@sap.com","IPSADMIN","t"),
    ("adminRM","Raj","Modi","N/A","N/A","N/A","","raj.modi@veritasprime.com","adminRM","t"),
    ("200001","Peter","Parker","Logistics (30000016)","Business Operations (40000015)","Atlanta (1710-2015)","Logistics Operative","","200001","t"),
    ("adminCA","Carlos","Arenas","N/A","N/A","N/A","","carlos.arenas@veritasprime.com","adminCA","t"),
    ("adminES","Elsa","Sanchez","N/A","N/A","N/A","","elsa.sanchez@veritasprime.com","adminES","t"),
    ("10000270","Spider","Contingent","Marketing (30000017)","Sales & Distribution (40000016)","Atlanta (1710-2015)","Sales Advisor","beinnennujubrou-8876@yopmail.com","10000270","t"),
    ("IPSADMIN_ay3hmungh","Cheryl","Lopez","N/A","N/A","N/A","","dummy@sap.com","IPSADMIN_ay3hmungh","t"),
    ("adminBG","Brigitte","Gutierrez","N/A","N/A","N/A","","brigitte.gutierrez@veritasprime.com","adminBG","t"),
    ("adminCS","Claver","Santos","N/A","N/A","N/A","","claver.santos@veritasprime.com","adminCS","t"),
    ("adminKT","Kayla","DePuy","N/A","N/A","N/A","","kayla.toelle+golden@veritasprime.com","adminKT","t"),
    ("adminMS","Manvitha Kiran","Singareddy","N/A","N/A","N/A","","manvitha.singareddy@veritasprime.com","adminMS","t"),
    ("adminVPProducts","Products","Admin","N/A","N/A","N/A","","shawna.sayar@veritasprime.com","adminVPProducts","t"),
    ("adminWN","Wilson","Naranjo","N/A","N/A","N/A","","wilsonanaranjo+golden@veritasprime.com","adminWN","t"),
    ("AdminSF","Veritas","Prime","N/A","N/A","N/A","","kayla.depuy+golden2@veritasprime.com","AdminSF","t"),
    ("200005","Natasha","Romanoff","Finance (30000006)","Finance (40000009)","Boston (1710-2017)","Finance Manager","","200005","t"),
    ("adminVP","Kayla","DePuy","N/A","N/A","N/A","","kayla.depuy+golden@veritasprime.com","adminVP","t"),
    ("adminDA","Danilo","Alarcon","N/A","N/A","N/A","","danilo.alarcon@veritasprime.com","adminDA","t"),
    ("adminVA","Victor","Alva","N/A","N/A","N/A","","victor.alva@veritasprime.com","adminVA","t"),
    ("adminSG","Sanjay","Gera","N/A","N/A","N/A","","sanjay.gera@veritasprime.com","adminSG","t"),
    ("10000251","Sal","Marullo","Logistics (30000016)","Business Operations (40000015)","Boston (1710-2017)","Sr. Logistics Operative","","10000251","t"),
    ("adminUB","Urmika","Banerjee","N/A","N/A","N/A","","urmika.banerjee@veritasprime.com","adminUB","t"),
    ("adminAL","Alex","Lindquist","N/A","N/A","N/A","","alex.lindquist@veritasprime.com","adminAL","t"),
    ("adminDS","Deepshikha","Singh","N/A","N/A","N/A","","deepshikha.singh@veritasprime.com","adminDS","t"),
    ("adminAH","Alberto","Hurtado","N/A","N/A","N/A","","alberto.hurtado@veritasprime.com","adminAH","t"),
    ("adminMN","Monica","Naranjo","N/A","N/A","N/A","","monica.naranjo@veritasprime.com","adminMN","t"),
    ("adminJB","Joe","Burkart","N/A","N/A","N/A","","joe.burkart+sandbox@veritasprime.com","adminJB","t"),
    ("200012","James","Howlett","Corporate (30000005)","Corporate Administrative (40000021)","Lima - Oficina Central (PE01_1001)","Director","","200012","t"),
    ("adminTG","Taylor","Compeau","N/A","N/A","N/A","","taylor.compeau+sandbox@veritasprime.com","adminTG","t"),
    ("adminWB","Willow","Boyce","N/A","N/A","N/A","","willow.boyce@veritasprime.com","adminWB","t"),
    ("adminJM","Jose","Machado","N/A","N/A","N/A","","jose.machado@veritasprime.com","adminJM","t"),
    ("VP_APIUSER_AL","VP API","AL","N/A","N/A","N/A","","dummy@successfactors.com","VP_APIUSER_AL","t"),
    ("VP_APIUSER_KK","VP API","KK","N/A","N/A","N/A","","dummy@successfactors.com","VP_APIUSER_KK","t"),
    ("VP_apiUser_CN","VP API","CN","N/A","N/A","N/A","","dummy@successfactors.com","VP_apiUser_CN","t"),
    ("VP_apiUser_DS","VP API","DS","N/A","N/A","N/A","","dummy@successfactors.com","VP_apiUser_DS","t"),
    ("api_vp","VP API","User","N/A","N/A","N/A","","dummy@successfactors.com","api_vp","t"),
    ("apiUserAD","VP API","AD","N/A","N/A","N/A","","dummy@successfactors.com","apiUserAD","t"),
    ("adminJP","Jennifer","Price","N/A","N/A","N/A","","jennifer.price@veritasprime.com","adminJP","t"),
    ("200022","Bucky","Barnes","Information Technology (30000009)","Information Technology (40000011)","Atlanta (1710-2015)","IT Director","","200022","t"),
    ("200007","Clint","Barton","Logistics (30000016)","Business Operations (40000015)","Atlanta (1710-2015)","Logistics Operative","noemail@company.com","200007","t"),
    ("10000253","Dragon","Fly","Corporate (30000005)","Corporate Administrative (40000021)","Alberta (CA01_1009)","Finance Manager","noemail@company.com","10000253","t"),
    ("200010","Carol","Danvers","Finance (30000006)","Finance (40000009)","Atlanta (1710-2015)","Finance Administrator","","200010","t"),
    ("10000265","Connor","Contingent","Manufacturing (30000001)","Production US (40000003)","Philadelphia (1710-2001)","CA CEO","noemail@company.com","10000265","t"),
    ("adminPD","Pooja","Danivasa","N/A","N/A","N/A","","pooja.danivasa+sandbox@veritasprime.com","adminPD","t"),
    ("200003","Steve","Rogers","Corporate (30000005)","Corporate Administrative (40000021)","Atlanta (1710-2015)","Chief Finance Officer Demo","","200003","t"),
    ("200004","Bruce","Banner","Finance (30000006)","Finance (40000009)","Atlanta (1710-2015)","Finance Manager","","200004","t"),
    ("adminEL","Edgar","Lopez","N/A","N/A","N/A","","edgar.lopez@veritasprime.com","adminEL","t"),
]

# Solo empleados con userId numerico
employees = [u for u in all_users if u[0].isdigit()]

wb = Workbook()
ws = wb.active
ws.title = "Empleados"

headers = ["userId","firstName","lastName","division","department","location","title","email","username","status"]

header_fill = PatternFill("solid", start_color="1F4E79")
header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = border

ws.row_dimensions[1].height = 30

data_font = Font(name="Arial", size=10)
alt_fill = PatternFill("solid", start_color="D6E4F0")

for row_idx, emp in enumerate(employees, 2):
    fill = alt_fill if row_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
    for col_idx, val in enumerate(emp, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = data_font
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(vertical="center")

col_widths = [20, 18, 18, 35, 38, 38, 28, 35, 20, 8]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{len(employees)+1}"

out = r"C:\Users\JoséMachado\Desktop\Proyectos\MCP SF\exports\empleados_VPMCSandbox.xlsx"

wb.save(out)
print(f"Guardado: {out} ({len(employees)} empleados)")
